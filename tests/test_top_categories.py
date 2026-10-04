"""'Top Categories' on the Dashboard of both spreadsheet reports."""
from __future__ import annotations

import pytest
from odf.opendocument import load as load_ods
from odf.table import Table, TableCell, TableRow
from odf.text import P
from openpyxl import load_workbook

from expense_tracker.ods import generate_ods
from expense_tracker.xlsx import generate_xlsx

from .conftest import add_transaction

# Twelve categories; the two that sort first by name are the smallest
AMOUNTS = {
    "Aaa tiny": 1.0, "Abb small": 2.0,
    "Cat03": 30.0, "Cat04": 40.0, "Cat05": 50.0, "Cat06": 60.0, "Cat07": 70.0,
    "Cat08": 80.0, "Cat09": 90.0, "Cat10": 100.0, "Utilities": 500.0, "Rent": 900.0,
}
TOP_TEN = [
    "Rent", "Utilities", "Cat10", "Cat09", "Cat08",
    "Cat07", "Cat06", "Cat05", "Cat04", "Cat03",
]


@pytest.fixture
def db(test_db):
    for category, amount in AMOUNTS.items():
        add_transaction(test_db, "2026-01-10", f"SHOP {category}", amount, category=category)
    # A refund: Cat10 nets 100 - 15 = 85, below Cat09
    add_transaction(test_db, "2026-01-12", "SHOP Cat10", 15.0, direction="in", category="Cat10")
    return test_db


def _between(labels: list[str], start: str, end: str) -> list[str]:
    i = labels.index(start)
    return [x for x in labels[i + 1 : labels.index(end, i)] if x]


def test_xlsx_lists_the_ten_largest_categories(db, tmp_path):
    path = tmp_path / "report.xlsx"
    generate_xlsx(db, tmp_path / "rules.csv", path, tmp_path / "notes.csv")
    ws = load_workbook(path)["Dashboard"]
    labels = [row[0].value or "" for row in ws.iter_rows(min_col=1, max_col=1)]

    expected = TOP_TEN[:2] + ["Cat09", "Cat10"] + TOP_TEN[4:]
    assert _between(labels, "Top Categories", "Top 10 Merchants") == expected


def test_ods_lists_the_ten_largest_categories(db, tmp_path):
    path = tmp_path / "report.ods"
    generate_ods(db, tmp_path / "rules.csv", path, tmp_path / "notes.csv")
    sheet = next(
        s
        for s in load_ods(str(path)).spreadsheet.getElementsByType(Table)
        if s.getAttribute("name") == "Dashboard"
    )
    labels = []
    for row in sheet.getElementsByType(TableRow):
        cells = row.getElementsByType(TableCell)
        labels.append("".join(str(p) for p in cells[0].getElementsByType(P)) if cells else "")

    expected = TOP_TEN[:2] + ["Cat09", "Cat10"] + TOP_TEN[4:]
    assert _between(labels, "Top Categories", "Top 10 Merchants") == expected
