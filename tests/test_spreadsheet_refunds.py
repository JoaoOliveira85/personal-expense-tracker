"""Refunds in spreadsheet reports: formulas must net them against spending.

No spreadsheet engine is available to evaluate the formulas, so these tests
pin the generated formula text.
"""
from __future__ import annotations

import sqlite3

import pytest
from odf.opendocument import load as load_ods
from odf.table import Table, TableCell, TableRow
from odf.text import P
from openpyxl import load_workbook

from expense_tracker.db import migrate_schema
from expense_tracker.ods import generate_ods
from expense_tracker.parser import _get_cleaning_patterns, reset_cleaning_cache
from expense_tracker.rules import categorize_transactions, load_rules
from expense_tracker.xlsx import generate_xlsx

N = 6  # last Data row: 5 sample transactions + header

XLSX_REFUNDS = (
    'Data!G:G,"in",Data!H:H,"<>",Data!H:H,"<>Income",Data!H:H,"<>Transfers"'
    ',Data!H:H,"<>Savings"'
)
ODS_SPEND = (
    f'(([.Data.G2:.Data.G{N}]="out")*([.Data.H2:.Data.H{N}]<>"Savings")'
    f'-([.Data.G2:.Data.G{N}]="in")*([.Data.H2:.Data.H{N}]<>"")'
    f'*([.Data.H2:.Data.H{N}]<>"Income")*([.Data.H2:.Data.H{N}]<>"Transfers")'
    f'*([.Data.H2:.Data.H{N}]<>"Savings"))'
    f'*[.Data.F2:.Data.F{N}]'
)


@pytest.fixture
def db(populated_db, rules_csv, noise_words, cleaning_patterns):
    """Sample DB, categorized, with the EXEMPLO payment turned into an
    Insurance refund (the only Insurance transaction) and the VODAFONE
    payment into a Savings deposit."""
    reset_cleaning_cache()
    _get_cleaning_patterns(noise_words, cleaning_patterns)
    conn = sqlite3.connect(str(populated_db))
    migrate_schema(conn)
    categorize_transactions(conn, load_rules(rules_csv))
    conn.execute(
        "UPDATE transactions SET direction = 'in', category = 'Insurance' "
        "WHERE description_raw LIKE '%EXEMPLO%'"
    )
    conn.execute(
        "UPDATE transactions SET category = 'Savings' "
        "WHERE description_raw LIKE '%VODAFONE%'"
    )
    conn.commit()
    conn.close()
    yield populated_db
    reset_cleaning_cache()


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------


@pytest.fixture
def workbook(db, rules_csv, tmp_path):
    path = tmp_path / "report.xlsx"
    generate_xlsx(db, rules_csv, path, tmp_path / "notes.csv")
    return load_workbook(path)


def _row_by_label(ws, label: str) -> int:
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == label:
            return row[0].row
    raise AssertionError(f"{label!r} not found in {ws.title}")


class TestXlsxRefunds:
    def test_category_total_nets_refunds(self, workbook):
        ws = workbook["Category Breakdown"]
        row = _row_by_label(ws, "Health")
        assert ws.cell(row=row, column=2).value == (
            '=SUMIFS(Data!F:F,Data!H:H,"Health",Data!G:G,"out")'
            '-SUMIFS(Data!F:F,Data!H:H,"Health",Data!G:G,"in")'
        )

    def test_category_with_only_refunds_is_listed(self, workbook):
        ws = workbook["Category Breakdown"]
        assert _row_by_label(ws, "Insurance")

    def test_dashboard_totals_net_refunds(self, workbook):
        ws = workbook["Dashboard"]
        income = ws.cell(row=_row_by_label(ws, "Total Income"), column=2).value
        expenses = ws.cell(row=_row_by_label(ws, "Total Expenses"), column=2).value
        assert income == (
            '=SUMIFS(Data!F:F,Data!G:G,"in",Data!H:H,"<>Savings")'
            f"-SUMIFS(Data!F:F,{XLSX_REFUNDS})"
        )
        assert expenses == (
            '=SUMIFS(Data!F:F,Data!G:G,"out",Data!H:H,"<>Savings")'
            f"-SUMIFS(Data!F:F,{XLSX_REFUNDS})"
        )

    def test_savings_is_not_a_spending_category(self, workbook):
        with pytest.raises(AssertionError):
            _row_by_label(workbook["Category Breakdown"], "Savings")

    def test_monthly_trend_nets_refunds(self, workbook):
        ws = workbook["Monthly Trend"]
        month = 'Data!B:B,"2026-01",'
        assert ws["B2"].value == (
            f'=SUMIFS(Data!F:F,{month}Data!G:G,"in",Data!H:H,"<>Savings")'
            f"-SUMIFS(Data!F:F,{month}{XLSX_REFUNDS})"
        )
        assert ws["C2"].value == (
            f'=SUMIFS(Data!F:F,{month}Data!G:G,"out",Data!H:H,"<>Savings")'
            f"-SUMIFS(Data!F:F,{month}{XLSX_REFUNDS})"
        )


# ---------------------------------------------------------------------------
# ODS
# ---------------------------------------------------------------------------


@pytest.fixture
def ods_doc(db, rules_csv, tmp_path):
    path = tmp_path / "report.ods"
    generate_ods(db, rules_csv, path, tmp_path / "notes.csv")
    return load_ods(str(path))


def _ods_row(doc, sheet_name: str, label: str) -> list:
    sheet = next(
        s
        for s in doc.spreadsheet.getElementsByType(Table)
        if s.getAttribute("name") == sheet_name
    )
    for row in sheet.getElementsByType(TableRow):
        cells = row.getElementsByType(TableCell)
        if cells and "".join(str(p) for p in cells[0].getElementsByType(P)) == label:
            return cells
    raise AssertionError(f"{label!r} not found in {sheet_name}")


class TestOdsRefunds:
    def test_category_total_nets_refunds(self, ods_doc):
        cells = _ods_row(ods_doc, "Category Breakdown", "Health")
        assert cells[1].getAttribute("formula") == (
            f'of:=SUMPRODUCT(([.Data.H2:.Data.H{N}]="Health")*{ODS_SPEND})'
        )

    def test_category_with_only_refunds_is_listed(self, ods_doc):
        assert _ods_row(ods_doc, "Category Breakdown", "Insurance")

    def test_savings_is_not_a_spending_category(self, ods_doc):
        with pytest.raises(AssertionError):
            _ods_row(ods_doc, "Category Breakdown", "Savings")

    def test_dashboard_total_expenses_nets_refunds(self, ods_doc):
        cells = _ods_row(ods_doc, "Dashboard", "Total Expenses")
        assert cells[1].getAttribute("formula") == f"of:=SUMPRODUCT({ODS_SPEND})"
