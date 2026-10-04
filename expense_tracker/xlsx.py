"""
XLSX report generation.

Generates the expense report in Excel format (.xlsx), which can also be
opened by Apple Numbers and Google Sheets.

This generates all the same sheets as the ODS version with working Excel formulas.
"""
from __future__ import annotations

import re
import sqlite3
from collections import defaultdict
from pathlib import Path

from .constants import (
    DEFAULT_DB,
    DEFAULT_DESC_NOTES,
    DATA_COLUMNS,
    DATA_HEADERS,
    NON_SPENDING_CATEGORIES,
    SAVINGS_CATEGORIES,
)
from .db import migrate_schema, fetch_all_transactions, counts_as_spending, is_savings
from .rules import load_rules

# SUMIFS criteria excluding savings, and selecting refunds (see db.py)
_NOT_SAVINGS = "".join(f',Data!H:H,"<>{c}"' for c in SAVINGS_CATEGORIES)
_REFUNDS = 'Data!G:G,"in",Data!H:H,"<>"' + "".join(
    f',Data!H:H,"<>{c}"' for c in NON_SPENDING_CATEGORIES + SAVINGS_CATEGORIES
)


def _spend(criteria: str = "") -> str:
    """SUMIFS spending net of refunds; `criteria` is e.g. 'Data!B:B,"2026-01",'."""
    return (
        f'SUMIFS(Data!F:F,{criteria}Data!G:G,"out"{_NOT_SAVINGS})'
        f"-SUMIFS(Data!F:F,{criteria}{_REFUNDS})"
    )


def _income(criteria: str = "") -> str:
    """SUMIFS incoming money that is neither a refund nor a savings withdrawal."""
    return (
        f'SUMIFS(Data!F:F,{criteria}Data!G:G,"in"{_NOT_SAVINGS})'
        f"-SUMIFS(Data!F:F,{criteria}{_REFUNDS})"
    )


def _criterion(name: str) -> str:
    """`name` as a quoted SUMIFS/COUNTIFS criterion that matches only itself.

    In a criterion * and ? are wildcards ("GLOVO*" would also add up every
    other merchant that starts with GLOVO), ~ escapes them, a leading =, <
    or > is a comparison, and a double quote ends the string.
    """
    escaped = re.sub(r"([~*?])", r"~\1", _printable(name))
    if escaped[:1] in ("=", "<", ">"):
        escaped = "=" + escaped
    return '"' + escaped.replace('"', '""') + '"'


def _category_spend(cat: str, criteria: str = "") -> str:
    """SUMIFS spending in one category net of its refunds."""
    criteria = f'{criteria}Data!H:H,{_criterion(cat)},'
    out = f'SUMIFS(Data!F:F,{criteria}Data!G:G,"out")'
    if cat in NON_SPENDING_CATEGORIES:
        return out
    return f'{out}-SUMIFS(Data!F:F,{criteria}Data!G:G,"in")'

# openpyxl imports - will fail gracefully if not installed
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def _printable(text: str) -> str:
    """`text` without the control characters an xlsx file cannot hold.

    openpyxl raises on them, so a single stray byte in a statement would
    cost the whole report. They become U+FFFD, as in the ODS report.
    """
    return ILLEGAL_CHARACTERS_RE.sub("\ufffd", text)


def _text_cell(ws, row: int, column: int, value):
    """Write data to a cell: statement text, a note, a category or rule name.

    openpyxl stores every string that starts with "=" as a formula, so a
    description such as "=1+1" would be evaluated by the spreadsheet.
    """
    if isinstance(value, str):
        value = _printable(value)
    cell = ws.cell(row=row, column=column, value=value)
    if cell.data_type == "f":
        cell.data_type = "s"
    return cell


def _load_description_notes(desc_notes_path: Path) -> dict[str, str]:
    """Load merchant notes from description-notes.csv."""
    import csv
    
    if not desc_notes_path.exists():
        return {}
    notes: dict[str, str] = {}
    with desc_notes_path.open("r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            key = row.get("description_clean", "").strip()
            val = row.get("merchant_note", "").strip()
            if key and val:
                notes[key] = val
    return notes


def generate_xlsx(
    db_path: Path,
    rules_path: Path,
    xlsx_path: Path,
    desc_notes_path: Path = DEFAULT_DESC_NOTES,
) -> None:
    """Generate an XLSX expense report with all analysis sheets."""
    if not OPENPYXL_AVAILABLE:
        print("Error: openpyxl is required for XLSX generation.")
        print("Install it with: pip install openpyxl")
        return
    
    conn = sqlite3.connect(str(db_path))
    try:
        migrate_schema(conn)
        transactions = fetch_all_transactions(conn)
    finally:
        conn.close()
    
    # Load rules and merchant notes
    rules = load_rules(rules_path) if rules_path.exists() else []
    desc_notes = _load_description_notes(desc_notes_path)
    for tx in transactions:
        tx["merchant_note"] = desc_notes.get(tx.get("description_clean", ""), "")
    
    # Create workbook
    wb = Workbook()
    
    # Define styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    section_fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
    categorized_fill = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
    uncategorized_fill = PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )
    
    styles = {
        'header_font': header_font,
        'header_fill': header_fill,
        'section_fill': section_fill,
        'categorized_fill': categorized_fill,
        'uncategorized_fill': uncategorized_fill,
        'thin_border': thin_border,
    }
    
    n = len(transactions) + 1  # Last data row (1-indexed, row 1 = header)
    
    # Create sheets
    ws_intro = wb.active
    ws_intro.title = "Intro"
    _write_intro_sheet(ws_intro, styles)
    
    ws_data = wb.create_sheet("Data")
    _write_data_sheet(ws_data, transactions, styles)
    
    ws_rules = wb.create_sheet("Rules")
    _write_rules_sheet(ws_rules, rules, styles)
    
    ws_dashboard = wb.create_sheet("Dashboard")
    _write_dashboard_sheet(ws_dashboard, transactions, styles, n)
    
    ws_monthly = wb.create_sheet("Monthly Summary")
    _write_monthly_summary_sheet(ws_monthly, transactions, styles, n)
    
    ws_trend = wb.create_sheet("Monthly Trend")
    _write_monthly_trend_sheet(ws_trend, transactions, styles, n)
    
    ws_category = wb.create_sheet("Category Breakdown")
    _write_category_breakdown_sheet(ws_category, transactions, styles, n)
    
    ws_subcategory = wb.create_sheet("Subcategory Breakdown")
    _write_subcategory_breakdown_sheet(ws_subcategory, transactions, styles, n)
    
    ws_tags = wb.create_sheet("Tags")
    _write_tags_sheet(ws_tags, transactions, styles)
    
    ws_recurring = wb.create_sheet("Recurring Merchants")
    _write_recurring_sheet(ws_recurring, transactions, styles)
    
    # Save
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(xlsx_path))
    
    print(f"XLSX report saved to {xlsx_path}")
    categorized = sum(1 for tx in transactions if tx.get("category"))
    total = len(transactions)
    print(
        f"  {total} transactions total, {categorized} categorized, "
        f"{total - categorized} uncategorized"
    )


def _write_intro_sheet(ws, styles: dict) -> None:
    """Write the intro/instructions sheet."""
    ws.column_dimensions['A'].width = 80
    
    intro_text = [
        "EXPENSE TRACKER",
        "",
        "This workbook contains your financial data and analysis.",
        "",
        "FORMULA RECALCULATION:",
        "If formulas show as 0 or #REF!, try:",
        "• Excel: Press Ctrl+Alt+F9 to recalculate all formulas",
        "• Numbers: Formulas should auto-calculate; if not, edit any cell and press Enter",
        "• Google Sheets: Formulas auto-calculate on import",
        "",
        "SHEETS:",
        "• Data - All transactions (you can edit Category, Subcategory, Notes here)",
        "• Rules - Categorization rules",
        "• Dashboard - Summary statistics",
        "• Monthly Summary - Spending by category per month",
        "• Monthly Trend - Income vs expenses over time",
        "• Category Breakdown - Total spending by category",
        "• Subcategory Breakdown - Detailed category analysis",
        "• Tags - Analysis of #tags in notes",
        "• Recurring Merchants - Merchants that appear regularly",
        "",
        "TIPS:",
        "• Use filters on the Data sheet to find specific transactions",
        "• Add #tags to the Notes column to track specific expenses",
        "• Categories and notes you edit will be preserved on re-import",
        "",
        "NOTE: This XLSX file works with Excel, Apple Numbers, and Google Sheets.",
        "For LibreOffice users, the ODS format (--format ods) is recommended.",
    ]
    
    for i, text in enumerate(intro_text, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        if i == 1:
            cell.font = Font(bold=True, size=16)
        elif text.startswith("SHEETS:") or text.startswith("TIPS:"):
            cell.font = Font(bold=True, size=12)


def _write_data_sheet(ws, transactions: list[dict], styles: dict) -> None:
    """Write the main data sheet with all transactions."""
    
    # Column widths
    col_widths = {
        'A': 12, 'B': 10, 'C': 6, 'D': 35, 'E': 30, 'F': 12, 'G': 8,
        'H': 15, 'I': 15, 'J': 12, 'K': 12, 'L': 8, 'M': 10, 'N': 15,
        'O': 12, 'P': 20, 'Q': 15, 'R': 25, 'S': 25,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    # Header row
    for col_idx, header in enumerate(DATA_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = Alignment(horizontal='center')
        cell.border = styles['thin_border']
    
    ws.freeze_panes = 'A2'
    
    # Data rows
    for row_idx, tx in enumerate(transactions, start=2):
        is_categorized = bool(tx.get("category"))
        fill = styles['categorized_fill'] if is_categorized else styles['uncategorized_fill']
        
        for col_idx, col_name in enumerate(DATA_COLUMNS, start=1):
            value = tx.get(col_name, "") or ""
            cell = _text_cell(ws, row_idx, col_idx, value)
            cell.fill = fill
            cell.border = styles['thin_border']
            
            if col_name in ("amount_abs", "balance") and isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal='right')
                cell.number_format = '#,##0.00'
    
    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DATA_HEADERS))}{len(transactions) + 1}"


def _write_rules_sheet(ws, rules: list[dict], styles: dict) -> None:
    """Write the rules sheet."""
    headers = ["Pattern", "Match Field", "Category", "Subcategory", "Payment Type"]
    
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.border = styles['thin_border']
        ws.column_dimensions[get_column_letter(col_idx)].width = 20
    
    for row_idx, rule in enumerate(rules, start=2):
        for col_idx, key in enumerate(["pattern", "match_field", "category", "subcategory", "payment_type"], start=1):
            cell = _text_cell(ws, row_idx, col_idx, rule.get(key, ""))
            cell.border = styles['thin_border']
    
    ws.freeze_panes = 'A2'


def _write_dashboard_sheet(ws, transactions: list[dict], styles: dict, n: int) -> None:
    """Write the dashboard summary sheet."""
    
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    
    n_months = len(set(tx["month"] for tx in transactions if tx["month"])) or 1
    
    row = 1
    
    # Title
    cell = ws.cell(row=row, column=1, value="DASHBOARD")
    cell.font = Font(bold=True, size=14, color="FFFFFF")
    cell.fill = styles['section_fill']
    ws.cell(row=row, column=2).fill = styles['section_fill']
    ws.cell(row=row, column=3).fill = styles['section_fill']
    row += 2
    
    # All-time summary header
    for col, header in enumerate(["All-Time Summary", "Amount", ""], start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
    row += 1
    
    # Total Income
    ws.cell(row=row, column=1, value="Total Income")
    ws.cell(row=row, column=2, value="=" + _income())
    ws.cell(row=row, column=2).number_format = '#,##0.00'
    row += 1
    
    # Total Expenses
    ws.cell(row=row, column=1, value="Total Expenses")
    ws.cell(row=row, column=2, value="=" + _spend())
    ws.cell(row=row, column=2).number_format = '#,##0.00'
    row += 1
    
    # Net Balance
    ws.cell(row=row, column=1, value="Net Balance (Income - Expenses)")
    ws.cell(row=row, column=2, value=f'=B4-B5')
    ws.cell(row=row, column=2).number_format = '#,##0.00'
    row += 1
    
    # Avg Monthly Spend
    ws.cell(row=row, column=1, value="Average Monthly Spend")
    ws.cell(row=row, column=2, value=f'=B5/{n_months}')
    ws.cell(row=row, column=2).number_format = '#,##0.00'
    row += 1
    
    # Total Transactions
    ws.cell(row=row, column=1, value="Total Transactions")
    ws.cell(row=row, column=2, value=len(transactions))
    row += 1
    
    # Uncategorized Count
    ws.cell(row=row, column=1, value="Uncategorized Transactions")
    ws.cell(row=row, column=2, value=f'=COUNTBLANK(Data!H2:H{n})')
    row += 2
    
    # Top Categories
    categories = sorted(set(
        tx["category"] for tx in transactions
        if tx.get("category") and counts_as_spending(tx)
    ))
    
    for col, header in enumerate(["Top Categories", "Total Spent", "# Transactions"], start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
    row += 1
    
    for cat in categories[:10]:
        _text_cell(ws, row, 1, cat)
        ws.cell(row=row, column=2, value="=" + _category_spend(cat))
        ws.cell(row=row, column=2).number_format = '#,##0.00'
        ws.cell(row=row, column=3, value=f'=COUNTIFS(Data!H:H,{_criterion(cat)},Data!G:G,"out")')
        row += 1
    
    row += 1
    
    # Top Merchants
    merchant_totals: dict[str, float] = {}
    for tx in transactions:
        if tx.get("direction") == "out" and not is_savings(tx):
            desc = tx.get("description_clean") or ""
            if desc:
                amt = float(tx.get("amount_abs") or 0)
                merchant_totals[desc] = merchant_totals.get(desc, 0) + amt
    
    top_merchants = sorted(merchant_totals.items(), key=lambda x: x[1], reverse=True)[:10]
    
    for col, header in enumerate(["Top 10 Merchants", "Total Spent", "# Transactions"], start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
    row += 1
    
    for merchant, total in top_merchants:
        _text_cell(ws, row, 1, merchant)
        match = _criterion(merchant)
        ws.cell(row=row, column=2, value=f'=SUMIFS(Data!F:F,Data!E:E,{match},Data!G:G,"out")')
        ws.cell(row=row, column=2).number_format = '#,##0.00'
        ws.cell(row=row, column=3, value=f'=COUNTIFS(Data!E:E,{match},Data!G:G,"out")')
        row += 1


def _write_monthly_summary_sheet(ws, transactions: list[dict], styles: dict, n: int) -> None:
    """Write monthly summary with category breakdown."""
    months = sorted(set(tx["month"] for tx in transactions if tx["month"]))
    categories = sorted(set(
        tx["category"] for tx in transactions
        if tx.get("category") and counts_as_spending(tx)
    ))
    
    if not categories:
        categories = ["(no categories yet)"]
    
    headers = ["Month"] + categories + ["Total"]
    
    # Headers
    for col_idx, header in enumerate(headers, start=1):
        cell = _text_cell(ws, 1, col_idx, header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.border = styles['thin_border']
        ws.column_dimensions[get_column_letter(col_idx)].width = 12
    
    ws.column_dimensions['A'].width = 10
    ws.freeze_panes = 'B2'
    
    # Data rows
    for row_idx, month in enumerate(months, start=2):
        ws.cell(row=row_idx, column=1, value=month).border = styles['thin_border']
        
        for col_idx, cat in enumerate(categories, start=2):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = "=" + _category_spend(cat, f'Data!B:B,"{month}",')
            cell.number_format = '#,##0.00'
            cell.border = styles['thin_border']
        
        # Total column
        last_cat_col = get_column_letter(len(categories) + 1)
        total_cell = ws.cell(row=row_idx, column=len(categories) + 2)
        total_cell.value = f'=SUM(B{row_idx}:{last_cat_col}{row_idx})'
        total_cell.number_format = '#,##0.00'
        total_cell.border = styles['thin_border']
    
    # Total row
    if months:
        total_row = len(months) + 2
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        for col_idx in range(2, len(categories) + 3):
            cell = ws.cell(row=total_row, column=col_idx)
            col_letter = get_column_letter(col_idx)
            cell.value = f'=SUM({col_letter}2:{col_letter}{total_row - 1})'
            cell.number_format = '#,##0.00'
            cell.font = Font(bold=True)


def _write_monthly_trend_sheet(ws, transactions: list[dict], styles: dict, n: int) -> None:
    """Write monthly trend (income vs expenses over time)."""
    months = sorted(set(tx["month"] for tx in transactions if tx["month"]))
    
    headers = ["Month", "Income", "Expenses", "Net", "Running Balance"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.border = styles['thin_border']
        ws.column_dimensions[get_column_letter(col_idx)].width = 15
    
    ws.freeze_panes = 'A2'
    
    for row_idx, month in enumerate(months, start=2):
        ws.cell(row=row_idx, column=1, value=month).border = styles['thin_border']
        
        # Income
        cell = ws.cell(row=row_idx, column=2)
        cell.value = "=" + _income(f'Data!B:B,"{month}",')
        cell.number_format = '#,##0.00'
        cell.border = styles['thin_border']
        
        # Expenses
        cell = ws.cell(row=row_idx, column=3)
        cell.value = "=" + _spend(f'Data!B:B,"{month}",')
        cell.number_format = '#,##0.00'
        cell.border = styles['thin_border']
        
        # Net
        cell = ws.cell(row=row_idx, column=4)
        cell.value = f'=B{row_idx}-C{row_idx}'
        cell.number_format = '#,##0.00'
        cell.border = styles['thin_border']
        
        # Running Balance
        cell = ws.cell(row=row_idx, column=5)
        if row_idx == 2:
            cell.value = f'=D{row_idx}'
        else:
            cell.value = f'=E{row_idx - 1}+D{row_idx}'
        cell.number_format = '#,##0.00'
        cell.border = styles['thin_border']
    
    # Total row
    if months:
        total_row = len(months) + 2
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        for col_idx, col_letter in enumerate(['B', 'C', 'D'], start=2):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.value = f'=SUM({col_letter}2:{col_letter}{total_row - 1})'
            cell.number_format = '#,##0.00'
            cell.font = Font(bold=True)


def _write_category_breakdown_sheet(ws, transactions: list[dict], styles: dict, n: int) -> None:
    """Write category breakdown analysis."""
    categories = sorted(set(
        tx["category"] for tx in transactions
        if tx.get("category") and counts_as_spending(tx)
    ))
    n_months = len(set(tx["month"] for tx in transactions if tx["month"])) or 1
    
    headers = ["Category", "Total Spent", "% of Total", "Avg / Month", "# Transactions"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.border = styles['thin_border']
        ws.column_dimensions[get_column_letter(col_idx)].width = 15
    
    ws.column_dimensions['A'].width = 20
    ws.freeze_panes = 'A2'
    
    for row_idx, cat in enumerate(categories, start=2):
        _text_cell(ws, row_idx, 1, cat).border = styles['thin_border']
        
        # Total Spent
        cell = ws.cell(row=row_idx, column=2)
        cell.value = "=" + _category_spend(cat)
        cell.number_format = '#,##0.00'
        cell.border = styles['thin_border']
        
        # % of Total
        cell = ws.cell(row=row_idx, column=3)
        cell.value = f"=B{row_idx}/({_spend()})*100"
        cell.number_format = '0.0"%"'
        cell.border = styles['thin_border']
        
        # Avg / Month
        cell = ws.cell(row=row_idx, column=4)
        cell.value = f'=B{row_idx}/{n_months}'
        cell.number_format = '#,##0.00'
        cell.border = styles['thin_border']
        
        # # Transactions
        cell = ws.cell(row=row_idx, column=5)
        cell.value = f'=COUNTIFS(Data!H:H,{_criterion(cat)},Data!G:G,"out")'
        cell.border = styles['thin_border']
    
    # Total row
    if categories:
        total_row = len(categories) + 2
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        for col_idx in range(2, 6):
            cell = ws.cell(row=total_row, column=col_idx)
            col_letter = get_column_letter(col_idx)
            cell.value = f'=SUM({col_letter}2:{col_letter}{total_row - 1})'
            cell.number_format = '#,##0.00'
            cell.font = Font(bold=True)


def _write_subcategory_breakdown_sheet(ws, transactions: list[dict], styles: dict, n: int) -> None:
    """Write subcategory breakdown analysis."""
    pairs = sorted(set(
        (tx["category"], tx.get("subcategory") or "")
        for tx in transactions
        if tx.get("category") and counts_as_spending(tx)
    ))
    
    headers = ["Category", "Subcategory", "Total Spent", "% of Category", "# Transactions"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.border = styles['thin_border']
        ws.column_dimensions[get_column_letter(col_idx)].width = 15
    
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 18
    ws.freeze_panes = 'A2'
    
    for row_idx, (cat, subcat) in enumerate(pairs, start=2):
        _text_cell(ws, row_idx, 1, cat).border = styles['thin_border']
        _text_cell(ws, row_idx, 2, subcat or "(none)").border = styles['thin_border']
        
        if subcat:
            # Total Spent
            cell = ws.cell(row=row_idx, column=3)
            cell.value = "=" + _category_spend(cat, f'Data!I:I,{_criterion(subcat)},')
            cell.number_format = '#,##0.00'
            cell.border = styles['thin_border']
            
            # % of Category
            cell = ws.cell(row=row_idx, column=4)
            cell.value = f"=C{row_idx}/({_category_spend(cat)})*100"
            cell.number_format = '0.0"%"'
            cell.border = styles['thin_border']
            
            # # Transactions
            cell = ws.cell(row=row_idx, column=5)
            cell.value = (
                f'=COUNTIFS(Data!H:H,{_criterion(cat)},'
                f'Data!I:I,{_criterion(subcat)},Data!G:G,"out")'
            )
            cell.border = styles['thin_border']
        else:
            # No subcategory - match empty
            cell = ws.cell(row=row_idx, column=3)
            cell.value = "=" + _category_spend(cat, 'Data!I:I,"",')
            cell.number_format = '#,##0.00'
            cell.border = styles['thin_border']
            
            cell = ws.cell(row=row_idx, column=4)
            cell.value = f"=C{row_idx}/({_category_spend(cat)})*100"
            cell.number_format = '0.0"%"'
            cell.border = styles['thin_border']
            
            cell = ws.cell(row=row_idx, column=5)
            cell.value = f'=COUNTIFS(Data!H:H,{_criterion(cat)},Data!I:I,"",Data!G:G,"out")'
            cell.border = styles['thin_border']


def _write_tags_sheet(ws, transactions: list[dict], styles: dict) -> None:
    """Write tags analysis sheet."""
    import re
    
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    
    # Extract tags from notes and merchant notes
    tag_pattern = re.compile(r'#\w+')
    tag_totals: dict[str, float] = defaultdict(float)
    tag_counts: dict[str, int] = defaultdict(int)
    
    for tx in transactions:
        if tx.get("direction") != "out":
            continue
        
        notes = (tx.get("notes") or "") + " " + (tx.get("merchant_note") or "")
        tags = tag_pattern.findall(notes)
        amount = float(tx.get("amount_abs") or 0)
        
        for tag in set(tags):  # Unique tags per transaction
            tag_totals[tag] += amount
            tag_counts[tag] += 1
    
    # Title
    cell = ws.cell(row=1, column=1, value="TAG ANALYSIS")
    cell.font = Font(bold=True, size=14, color="FFFFFF")
    cell.fill = styles['section_fill']
    ws.cell(row=1, column=2).fill = styles['section_fill']
    ws.cell(row=1, column=3).fill = styles['section_fill']
    
    row = 3
    
    if not tag_totals:
        ws.cell(row=row, column=1, value="No #tags found in Notes or Merchant Notes.")
        ws.cell(row=row + 2, column=1, value="Tip: Add #tags like #groceries, #subscription, #gift to track specific expenses.")
        return
    
    # Headers
    for col, header in enumerate(["Tag", "Total Spent", "# Transactions"], start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.border = styles['thin_border']
    row += 1
    
    # Sorted by total descending
    for tag, total in sorted(tag_totals.items(), key=lambda x: -x[1]):
        ws.cell(row=row, column=1, value=tag).border = styles['thin_border']
        
        cell = ws.cell(row=row, column=2, value=total)
        cell.number_format = '#,##0.00'
        cell.border = styles['thin_border']
        
        ws.cell(row=row, column=3, value=tag_counts[tag]).border = styles['thin_border']
        row += 1


def _write_recurring_sheet(ws, transactions: list[dict], styles: dict) -> None:
    """Write recurring merchants analysis."""
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    
    # Find merchants that appear in multiple months
    merchant_months: dict[str, set[str]] = defaultdict(set)
    merchant_totals: dict[str, float] = defaultdict(float)
    merchant_counts: dict[str, int] = defaultdict(int)
    
    for tx in transactions:
        if tx.get("direction") != "out":
            continue
        
        desc = tx.get("description_clean") or ""
        month = tx.get("month") or ""
        amount = float(tx.get("amount_abs") or 0)
        
        if desc and month:
            merchant_months[desc].add(month)
            merchant_totals[desc] += amount
            merchant_counts[desc] += 1
    
    # Filter to merchants appearing in 2+ months
    recurring = {
        m: months for m, months in merchant_months.items()
        if len(months) >= 2
    }
    
    # Title
    cell = ws.cell(row=1, column=1, value="RECURRING MERCHANTS")
    cell.font = Font(bold=True, size=14, color="FFFFFF")
    cell.fill = styles['section_fill']
    for col in range(2, 5):
        ws.cell(row=1, column=col).fill = styles['section_fill']
    
    row = 3
    
    if not recurring:
        ws.cell(row=row, column=1, value="No recurring merchants found (merchants appearing in 2+ months).")
        return
    
    # Headers
    for col, header in enumerate(["Merchant", "# Months", "Total Spent", "Avg / Month"], start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.border = styles['thin_border']
    row += 1
    
    # Sorted by number of months descending, then by total
    sorted_merchants = sorted(
        recurring.items(),
        key=lambda x: (-len(x[1]), -merchant_totals[x[0]])
    )
    
    for merchant, months in sorted_merchants[:30]:  # Top 30
        _text_cell(ws, row, 1, merchant).border = styles['thin_border']
        ws.cell(row=row, column=2, value=len(months)).border = styles['thin_border']
        
        cell = ws.cell(row=row, column=3, value=merchant_totals[merchant])
        cell.number_format = '#,##0.00'
        cell.border = styles['thin_border']
        
        cell = ws.cell(row=row, column=4, value=merchant_totals[merchant] / len(months))
        cell.number_format = '#,##0.00'
        cell.border = styles['thin_border']
        
        row += 1


